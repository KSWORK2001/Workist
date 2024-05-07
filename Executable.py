import re
import docx
import tkinter as tk
from tkinter import filedialog
import spacy

def parse_docx_resume(resume_file):
    try:
        doc = docx.Document(resume_file)
        text = []
        for paragraph in doc.paragraphs:
            text.append(paragraph.text)

        resume_text = '\n'.join(text)

        # Now you can proceed with your regular expression-based parsing or any other method
        resume_info = extract_info_from_text(resume_text)
        return resume_info
    except Exception as e:
        print("Error:", e)
        return None

def extract_info_from_text(resume_text):
    try:
        name_match = re.search(r'^\s*([\w\s]+)$', resume_text, re.MULTILINE)
        name = name_match.group(1) if name_match else None

        # Extract phone number
        phone_match = re.search(r'(?<=\|)\s*\(?(\d{3})\)?[\s.-]?(\d{3})[\s.-]?(\d{4})', resume_text)
        phone = f"{''.join(phone_match.groups())}" if phone_match else None

        # Extract email
        email_match = re.search(r'[\w.-]+@[a-zA-Z.-]+\.[a-zA-Z]{2,4}', resume_text)
        email = email_match.group() if email_match else None

        # Extract education
        education_match = re.search(r'Education:(.*?)Expected Graduation:', resume_text, re.DOTALL)
        education = education_match.group(1).strip() if education_match else None
        
        # Extract graduation date
        gradDate_match = re.search(r'Expected Graduation:(.*?)(B\.S\.|B\.A\.|Bachelors|Bachelor\'s|M\.S\.|Master\'s|Associate\'s|Diploma)', resume_text, re.DOTALL)
        gradDate = gradDate_match.group(1).strip() if gradDate_match else None
        
        # Extract major
        major_match = re.search(r'(?:B\.S\.|B\.A\.|Bachelors|Bachelor\'s|M\.S\.|Master\'s|Associate\'s|Diploma)(?:\s+in|\'s in|at)?\s+(.*?)(?:GPA:|Activities/Honors:|Work Experience:)', resume_text, re.DOTALL)
        major = major_match.group(1).strip() if major_match else None
        
        # Extract degree
        degree_match = re.search(r'(B\.S\.|B\.A\.|Bachelors|Bachelor\'s|M\.S\.|Master\'s|Associate\'s|Diploma)(.*?)\b(?:in|at)\b', resume_text, re.DOTALL)
        degree = degree_match.group(1).strip() if degree_match else None
                
        # Extract major
        major_match = re.search(r'(?:B\.S\.|B\.A\.|Bachelors|Bachelor\'s|M\.S\.|Master\'s|Associate\'s|Diploma)(?:\s+in|\'s in|at)?\s+(.*?)(?:GPA:|Activities/Honors:|Work Experience:)', resume_text, re.DOTALL)
        major = major_match.group(1).strip() if major_match else None
        
        # Extract GPA
        gpa_match = re.search(r'GPA:(.*?)Honors/Activities:', resume_text, re.DOTALL)
        gpa = gpa_match.group(1).strip() if gpa_match else None
        
        # Extract activities/honors
        activities_match = re.search(r'Honors/Activities:(.*?)Work Experience:', resume_text, re.DOTALL)
        activities = activities_match.group(1).strip() if activities_match else None

        # Extract work experience
        work_experience = {}
        work_exp_sections = re.findall(r'(.*?)\s*(\d{4}-\d{4})\s*\n(.*?)\n(?=.*?(?:Leadership Roles|Leadership Roles and Personal Projects))', resume_text, re.DOTALL)
        for section in work_exp_sections:
            heading = section[0].strip()
            date_range = section[1].strip()
            job_description = section[2].strip()
            work_experience[heading] = {'Date Range': date_range, 'Job Description': job_description}

        # Extract leadership roles and personal projects
        leadership_projects_match = re.search(r'(Leadership Roles and Personal Projects:|Projects:|Leadership Roles:)(.*?)Certifications', resume_text, re.DOTALL)
        leadership_projects = leadership_projects_match.group(1).strip() if leadership_projects_match else None

        # Extract certifications
        certifications_match = re.search(r'Certifications:(.*?)(Technical Skills|Skills)', resume_text, re.DOTALL)
        certifications = certifications_match.group(1).strip() if certifications_match else None

        # Extract technical skills
        technical_skills_match = re.search(r'(Technical Skills:|Skills:)(.*)', resume_text, re.DOTALL)
        technical_skills = technical_skills_match.group(1).strip() if technical_skills_match else None

        return {
            'Name': name,
            'Phone': phone,
            'Email': email,
            'Education': education,
            'Graduation Date' : gradDate,
            'GPA' : gpa,
            'Activities/Honors': activities,
            'Major': major,
            'Degree': degree,
            'Work Experience': work_experience,
            'Leadership Projects': leadership_projects,
            'Certifications': certifications,
            'Skills': technical_skills
        }
    except AttributeError:
        print("Error: Failed to extract information from resume text.")
        return None

def browse_file():
    root = tk.Tk()
    root.withdraw()  # Hide the main window

    file_path = filedialog.askopenfilename()  # Open the file dialog
    return file_path

def display_welcome_gui():
    root = tk.Tk()
    root.title("Resume Parser")

    # Set window size and position it in the middle of the screen
    window_width = 600
    window_height = 400
    screen_width = root.winfo_screenwidth()
    screen_height = root.winfo_screenheight()
    x_position = (screen_width - window_width) // 2
    y_position = (screen_height - window_height) // 2
    root.geometry(f"{window_width}x{window_height}+{x_position}+{y_position}")

    # Set background color
    root.configure(bg="#f0f0f0")

    # Create a frame for the content
    content_frame = tk.Frame(root, bg="#ffffff", pady=20)
    content_frame.pack(expand=True, fill="both", padx=20, pady=(50, 20))

    # Welcome label
    welcome_label = tk.Label(content_frame, text="Cover AI", font=("Helvetica", 24, "bold"), bg="#ffffff", pady=20)
    welcome_label.pack()

    # Instruction label
    instruction_label = tk.Label(content_frame, text="Please select your resume file:", font=("Helvetica", 14), bg="#ffffff")
    instruction_label.pack()

    # Browse button
    browse_button = tk.Button(content_frame, text="Browse", command=select_resume, font=("Helvetica", 12), bg="#4caf50", fg="white", padx=20, pady=10)
    browse_button.pack(pady=20)

    root.mainloop()

def select_resume():
    resume_path = browse_file()
    if resume_path:
        resume_info = parse_docx_resume(resume_path)
        if resume_info:
            display_resume_info(resume_info)
        else:
            display_error_message("Failed to parse resume.")
    else:
        display_error_message("No file selected.")

def display_resume_info(resume_info):
    info_window = tk.Toplevel()
    info_window.title("Resume Information")

    # Set background color
    info_window.configure(bg="#f0f0f0")

    # Create a frame for the content
    content_frame = tk.Frame(info_window, bg="#ffffff", padx=20, pady=20)
    content_frame.pack(expand=True, fill="both")

    # Title label
    title_label = tk.Label(content_frame, text="Resume Information", font=("Helvetica", 18, "bold"), bg="#ffffff", pady=10)
    title_label.pack()

    # Company name entry
    company_frame = tk.Frame(content_frame, bg="#ffffff")
    company_frame.pack(fill="x", pady=5)

    company_label = tk.Label(company_frame, text="Company:", font=("Helvetica", 12, "bold"), bg="#ffffff")
    company_label.pack(side="left", padx=5)

    company_entry = tk.Entry(company_frame, font=("Helvetica", 12), bg="#f0f0f0")
    company_entry.pack(side="left", fill="x", expand=True)

    # Years of experience entry
    experience_frame = tk.Frame(content_frame, bg="#ffffff")
    experience_frame.pack(fill="x", pady=5)

    experience_label = tk.Label(experience_frame, text="Years of Experience:", font=("Helvetica", 12, "bold"), bg="#ffffff")
    experience_label.pack(side="left", padx=5)

    experience_entry = tk.Entry(experience_frame, font=("Helvetica", 12), bg="#f0f0f0")
    experience_entry.pack(side="left", fill="x", expand=True)
    
    # Field of interest dropdown
    field_of_interest_frame = tk.Frame(content_frame, bg="#ffffff")
    field_of_interest_frame.pack(fill="x", pady=5)

    field_of_interest_label = tk.Label(field_of_interest_frame, text="Field of Interest:", font=("Helvetica", 12, "bold"), bg="#ffffff")
    field_of_interest_label.pack(side="left", padx=5)

    field_options = ["Front End Dev", "Back End Dev", "Full Stack Dev", "Data Analyst", "Cloud Architect", "Project Manager"]
    field_of_interest_var = tk.StringVar(info_window)
    field_of_interest_var.set(field_options[0])
    field_of_interest_dropdown = tk.OptionMenu(field_of_interest_frame, field_of_interest_var, *field_options)
    field_of_interest_dropdown.pack(side="left", fill="x", expand=True)


    # Create a label for each field and its value
    entry_widgets = {}
    for key, value in resume_info.items():
        if key != "Work Experience":
            label_frame = tk.Frame(content_frame, bg="#ffffff")
            label_frame.pack(fill="x", pady=5)

            label_key = tk.Label(label_frame, text=key + ":", font=("Helvetica", 12, "bold"), bg="#ffffff")
            label_key.pack(side="left", padx=5)

            # For other fields, display editable entry widgets
            entry = tk.Entry(label_frame, font=("Helvetica", 12), bg="#f0f0f0")
            entry.pack(side="left", fill="x", expand=True)
            entry.insert(0, str(value))
            entry_widgets[key] = entry
    
     # Add entry fields for work experience
    work_exp_frame = tk.Frame(content_frame, bg="#ffffff")
    work_exp_frame.pack(fill="x", pady=5)

    work_exp_label = tk.Label(work_exp_frame, text="Work Experience:", font=("Helvetica", 12, "bold"), bg="#ffffff")
    work_exp_label.pack(side="left", padx=5)

    work_exp_entry = tk.Text(work_exp_frame, font=("Helvetica", 12), bg="#f0f0f0", height=5)
    work_exp_entry.pack(side="left", fill="both", expand=True)

     # Add entry field for previous company
    previous_company_frame = tk.Frame(content_frame, bg="#ffffff")
    previous_company_frame.pack(fill="x", pady=5)

    previous_company_label = tk.Label(previous_company_frame, text="Previous Company:", font=("Helvetica", 12, "bold"), bg="#ffffff")
    previous_company_label.pack(side="left", padx=5)

    previous_company_entry = tk.Entry(previous_company_frame, font=("Helvetica", 12), bg="#f0f0f0")
    previous_company_entry.pack(side="left", fill="x", expand=True)  

    # Generate cover letter button
    generate_cover_letter_button = tk.Button(content_frame, text="Generate Cover Letter", font=("Helvetica", 12), command=lambda: generate_cover_letter(resume_info, company_entry.get(), experience_entry.get(), field_of_interest_var.get(), previous_company_entry.get()))
    generate_cover_letter_button.pack(side="bottom", pady=10)

    info_window.mainloop()

        
def generate_cover_letter(resume_info, company, experience, field_of_interest, previous_company):
    
    # Getting info from resume
    name = resume_info['Name']
    phone = resume_info['Phone']
    email = resume_info['Email']
    education = resume_info['Education']
    gradDate = resume_info['Graduation Date']
    gpa = resume_info['GPA']
    activities = resume_info['Activities/Honors']
    major = resume_info['Major']
    degree = resume_info['Degree']
    workex = resume_info['Work Experience']
    leadership_projects = resume_info['Leadership Projects']
    certificates = resume_info['Certifications']
    skills = resume_info['Skills']

    # Initialize the cover letter template
    cover_letter_template = ""

    # Select cover letter template based on field of interest
    if field_of_interest == "Front End Dev":
        cover_letter_template = """
        Dear Hiring Manager,

        I am writing to express my interest in the Front End Developer position at {company}. With {experience} years of experience in web development, I am confident in my ability to contribute effectively to your team.

        Based on my review of the job description, I believe that my skills in HTML, CSS, and JavaScript align well with the requirements of the role. In my previous role at {previous_company}, I successfully led the development of several responsive web applications, demonstrating my ability to deliver high-quality front end solutions.

        I am particularly excited about the opportunity to work at {company} because of your commitment to innovation and user experience. I am eager to bring my expertise in front end development to help {company} achieve its goals.

        Thank you for considering my application. I look forward to the opportunity to discuss how my skills and experiences align with the needs of your team.

        Sincerely,
        {name}
        """
    elif field_of_interest == "Back End Dev":
        cover_letter_template = """
        Dear Hiring Manager,

        I am writing to apply for the Back End Developer position at {company}. With {experience} years of experience in software development, I am confident in my ability to contribute to your team and support your backend infrastructure.

        My experience with backend technologies such as Python, Java, and SQL, combined with my strong problem-solving skills, make me well-suited for the challenges of this role. In my previous role at {previous_company}, I successfully developed and maintained scalable backend systems, ensuring optimal performance and reliability.

        I am particularly impressed by {company}'s innovative approach to backend development and the opportunity to work with cutting-edge technologies. I am eager to leverage my skills and expertise to help {company} achieve its objectives.

        Thank you for considering my application. I am excited about the possibility of joining your team and contributing to your success.

        Sincerely,
        {name}
        """
    elif field_of_interest == "Full Stack Dev":
        cover_letter_template = """
        Dear Hiring Manager,

        I am writing to express my interest in the Full Stack Developer position at {company}. With {experience} years of experience in software development, I have a strong foundation in both front end and back end technologies.

        My proficiency in HTML, CSS, JavaScript, Python, and SQL allows me to develop end-to-end solutions that meet the needs of modern web applications. In my previous role at {previous_company}, I successfully led the development of several full stack projects, demonstrating my ability to deliver high-quality software solutions.

        I am particularly excited about the opportunity to work at {company} because of your commitment to innovation and technology. I am eager to contribute my skills and expertise to help {company} achieve its goals.

        Thank you for considering my application. I look forward to the opportunity to discuss how I can contribute to your team.

        Sincerely,
        {name}
        """
    elif field_of_interest == "Data Analyst":
        cover_letter_template = """
        Dear Hiring Manager,

        I am writing to apply for the Data Analyst position at {company}. With {experience} years of experience in data analysis, I am confident in my ability to leverage data to drive informed decision-making and support your business objectives.

        My expertise in data mining, statistical analysis, and data visualization tools such as Python, R, and Tableau, make me well-suited for the challenges of this role. In my previous role at {previous_company}, I successfully analyzed large datasets to uncover insights and trends, contributing to strategic business initiatives.

        I am particularly impressed by {company}'s commitment to data-driven decision-making and the opportunity to work with a talented team of professionals. I am excited about the possibility of applying my skills and expertise to help {company} achieve its goals.

        Thank you for considering my application. I look forward to the opportunity to discuss how I can contribute to your team.

        Sincerely,
        {name}
        """
    elif field_of_interest == "Cloud Architect":
        cover_letter_template = """
        Dear Hiring Manager,

        I am writing to express my interest in the Cloud Architect position at {company}. With {experience} years of experience in cloud computing and infrastructure design, I am confident in my ability to architect scalable and reliable cloud solutions.

        My expertise in cloud platforms such as AWS, Azure, and Google Cloud, combined with my strong understanding of networking and security principles, make me well-suited for the challenges of this role. In my previous role at {previous_company}, I successfully designed and implemented cloud solutions that improved efficiency and reduced costs.

        I am particularly excited about the opportunity to work at {company} because of your commitment to innovation and technology. I am eager to contribute my skills and expertise to help {company} leverage the power of the cloud to achieve its objectives.

        Thank you for considering my application. I look forward to the opportunity to discuss how I can contribute to your team.

        Sincerely,
        {name}
        """
    elif field_of_interest == "Project Manager":
        cover_letter_template = """
        Dear Hiring Manager,

        I am writing to apply for the Project Manager position at {company}. With {experience} years of experience in project management, I am confident in my ability to lead cross-functional teams and deliver projects on time and within budget.

        My strong organizational skills, attention to detail, and effective communication abilities have enabled me to successfully manage projects from initiation to completion. In my previous role at {previous_company}, I successfully led several high-impact projects, resulting in increased efficiency and customer satisfaction.

        I am particularly impressed by {company}'s commitment to excellence and the opportunity to work with a talented team of professionals. I am eager to bring my skills and expertise to help {company} achieve its strategic objectives.

        Thank you for considering my application. I look forward to the opportunity to discuss how I can contribute to your team.

        Sincerely,
        {name}
        """
    else:
        # Handle the case if the field of interest is not recognized
        print("Error: Invalid field of interest.")

    # Fill in placeholders in the cover letter template with information from the resume
    cover_letter_text = cover_letter_template.format(company=company, experience=experience, name = name, previous_company = previous_company)

    # Display the generated cover letter
    cover_letter_window = tk.Toplevel()
    cover_letter_window.title("Cover Letter")

    # Create a text widget to display the cover letter
    cover_letter_text_widget = tk.Text(cover_letter_window, wrap="word", font=("Helvetica", 12))
    cover_letter_text_widget.insert("1.0", cover_letter_text)
    cover_letter_text_widget.pack(expand=True, fill="both", padx=20, pady=20)

    cover_letter_window.mainloop()

 

def display_error_message(message):
    error_window = tk.Toplevel()
    error_window.title("Error")

    error_label = tk.Label(error_window, text=message, font=("Helvetica", 12))
    error_label.pack(padx=20, pady=10)

display_welcome_gui()
